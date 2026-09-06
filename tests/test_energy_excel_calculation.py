"""Excel-conversion energy schedule calculation tests."""

from io import BytesIO
from datetime import date

from fastapi import FastAPI
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from api.routers.energy_schedule import router
from api.routers.reports import router as reports_router
from api.schemas.energy_excel_calculation import ConsumptionEntry
from api.services.energy_excel_calculation_service import (
    CalculationAssumptions,
    DailyCalculationInput,
    MarketInput,
    build_calculation_trace,
    calculate_daily,
    calculate_savings_sheet,
    calculate_slot_wise_consolidate,
    compare_daily,
    list_consumption,
    run_calculation,
    upsert_consumption,
)
from api.services.energy_excel_report_service import (
    generate_energy_excel_calculation_report,
    generate_energy_pdf_calculation_report,
)
from database.config import Base
from database.models import Client, EnergyScheduleConsumption, MonthlyCalculation, Portfolio


def test_day_01_parity_matches_sample_workbook_outputs() -> None:
    assumptions = CalculationAssumptions()
    inputs = DailyCalculationInput(
        portfolio_id=1,
        trading_date=date(2026, 8, 1),
        consumption=ConsumptionEntry(
            portfolio_id=1,
            consumption_date=date(2026, 8, 1),
            c1_kwh=1500,
            c2_kwh=9300,
            c4_kwh=18300,
            c5_kwh=5150,
            base_tariff_per_unit=7.5,
        ),
        markets={
            "gdam": MarketInput(
                market="gdam",
                purchased_mwh=2.55,
                iex_payout=6850.77,
                ctu_loss_pct=0.0379,
                ctu_charge=122.94,
                nldc_fee=5.64,
                bucket_mwh={"c1": 2.55, "c2": 0.0, "c4": 0.0, "c5": 0.0},
            ),
            "dam": MarketInput(market="dam"),
            "rtm": MarketInput(market="rtm"),
        },
    )

    result = calculate_daily(inputs, "parity", assumptions)

    assert result.b44_iex_price == pytest_approx(16224.506875251094)
    assert result.b45_iex_price_per_unit == pytest_approx(6.906321519136511)
    assert result.b46_eb_price == pytest_approx(23125.188873804607)
    assert result.b47_eb_price_per_unit == pytest_approx(9.84375)


def test_corrected_mode_avoids_spreadsheet_division_errors() -> None:
    inputs = DailyCalculationInput(
        portfolio_id=1,
        trading_date=date(2026, 8, 4),
        consumption=None,
        markets={
            "gdam": MarketInput(market="gdam"),
            "dam": MarketInput(market="dam"),
            "rtm": MarketInput(market="rtm"),
        },
    )

    parity = calculate_daily(inputs, "parity")
    corrected = calculate_daily(inputs, "corrected")
    comparison = compare_daily(inputs)

    assert parity.b45_iex_price_per_unit == "#DIV/0!"
    assert parity.b47_eb_price_per_unit == "#DIV/0!"
    assert corrected.b45_iex_price_per_unit is None
    assert corrected.b47_eb_price_per_unit is None
    assert "consumption" in corrected.missing_inputs
    assert comparison.differences["b45_iex_price_per_unit"] is None


def test_savings_sheet_rolls_up_sample_workbook_days() -> None:
    assumptions = CalculationAssumptions()
    comparisons = [compare_daily(item, assumptions) for item in sample_workbook_inputs()]

    sheet = calculate_savings_sheet(comparisons, use_corrected=False)
    corrected_sheet = calculate_savings_sheet(comparisons)

    assert len(sheet.rows) == 3
    assert sheet.rows[0].iex_price == pytest_approx(16224.506875251094)
    assert sheet.rows[0].equivalent_eb_price == pytest_approx(23125.188873804607)
    assert sheet.rows[0].total_cost_saving == pytest_approx(6900.681998553513)
    assert sheet.rows[0].line_loss_pct == pytest_approx(0.0379)
    assert sheet.totals["iex_price"] == pytest_approx(56016.0146258943)
    assert sheet.totals["equivalent_eb_price"] == pytest_approx(78211.34156493431)
    assert sheet.totals["total_cost_saving"] == pytest_approx(22195.32693904001)
    assert sheet.totals["iex_price_per_unit_average"] == pytest_approx(7.677022442318424)
    assert sheet.totals["equivalent_eb_price_per_unit_average"] == pytest_approx(10.71875)
    assert corrected_sheet.rows[0].iex_price == pytest_approx(16214.707437751096)


def test_slot_wise_consolidate_matches_sample_workbook_slots() -> None:
    sheet = calculate_slot_wise_consolidate(sample_workbook_inputs())
    rows = {row.slot: row for row in sheet.rows}

    assert rows["C1"].consumption_kwh == pytest_approx(5350)
    assert rows["C1"].iex_delivered_kwh == pytest_approx(7286.919895255269)
    assert rows["C1"].balance_kwh == pytest_approx(-1936.9198952552688)
    assert rows["C2"].consumption_kwh == pytest_approx(29950)
    assert rows["C2"].iex_delivered_kwh == pytest_approx(0)
    assert rows["C4"].consumption_kwh == pytest_approx(57900)
    assert rows["C5"].consumption_kwh == pytest_approx(12600)


def test_consumption_upsert_and_month_calculation_persist_separately() -> None:
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)
    db = sessionmaker(bind=engine)()
    try:
        client = Client(entity_id="TEST", entity_name="Test Client")
        db.add(client)
        db.flush()
        portfolio = Portfolio(client_id=client.id, portfolio_code="TEST-PF", portfolio_name="Test Portfolio")
        db.add(portfolio)
        db.commit()

        record = ConsumptionEntry(
            portfolio_id=portfolio.id,
            consumption_date=date(2026, 8, 1),
            c1_kwh=1,
            c2_kwh=2,
            c4_kwh=3,
            c5_kwh=4,
            base_tariff_per_unit=8.5,
        )
        upsert_consumption(db, [record])
        record.c1_kwh = 10
        upsert_consumption(db, [record])

        rows = list_consumption(db, portfolio.id, date(2026, 8, 1), date(2026, 8, 1))
        assert len(rows) == 1
        assert rows[0].c1_kwh == 10
        assert db.query(EnergyScheduleConsumption).count() == 1

        results = run_calculation(db, portfolio.id, 2026, 8, "compare", day=1)
        assert len(results) == 1
        assert db.query(MonthlyCalculation).count() == 1
        assert db.query(MonthlyCalculation).first().calculation_type.endswith(":compare")

        month_results = run_calculation(db, portfolio.id, 2026, 8, "compare")
        assert len(month_results) == 31
        assert db.query(MonthlyCalculation).count() == 32
        summary = db.query(MonthlyCalculation).filter(
            MonthlyCalculation.calculation_type == "excel-conversion-v1:monthly:compare"
        ).one()
        assert summary.day == 0
        assert summary.calculation_data["calculation_version"] == "excel-conversion-v1"
        assert summary.calculation_data["calculation_mode"] == "compare"
        assert summary.calculation_data["completion"]["days_total"] == 31
        assert "source_files" in summary.calculation_data["audit"]
        assert summary.calculation_data["audit"]["consumption"][0]["source"] == "manual"
        assert "savings_sheet" in summary.calculation_data
        assert "slot_wise_consolidate" in summary.calculation_data

        trace = build_calculation_trace(db, portfolio.id, 2026, 8, day=1)
        assert len(trace["daily_outputs"]) == 1
        assert trace["consumption_inputs"][0]["present"] is True
        assert "savings_sheet" in trace
        assert "slot_wise_consolidate" in trace
    finally:
        db.close()
        Base.metadata.drop_all(engine)


def test_saved_monthly_summary_generates_clean_excel_report() -> None:
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)
    db = sessionmaker(bind=engine)()
    try:
        client = Client(entity_id="TEST", entity_name="Test Client")
        db.add(client)
        db.flush()
        portfolio = Portfolio(client_id=client.id, portfolio_code="TEST-PF", portfolio_name="Test Portfolio")
        db.add(portfolio)
        db.commit()

        run_calculation(db, portfolio.id, 2026, 8, "compare")
        summary = db.query(MonthlyCalculation).filter(
            MonthlyCalculation.calculation_type == "excel-conversion-v1:monthly:compare"
        ).one()

        buffer = generate_energy_excel_calculation_report(summary, portfolio)
        workbook = load_workbook(filename=BytesIO(buffer.getvalue()), data_only=True)

        assert workbook.sheetnames == ["Summary", "Daily Results", "Slot Wise Consolidate", "Assumptions"]
        assert workbook["Summary"]["A1"].value == "Energy Schedule Savings Report"
        assert workbook["Summary"]["A8"].value == "Completed Days"
        assert workbook["Daily Results"]["A3"].value == "Date"
        assert workbook["Daily Results"]["A4"].value == "2026-08-01"
        assert workbook["Slot Wise Consolidate"]["A3"].value == "Slot"
        assert "parity" not in {sheet.lower() for sheet in workbook.sheetnames}

        pdf_buffer = generate_energy_pdf_calculation_report(summary, portfolio)
        assert pdf_buffer.getvalue().startswith(b"%PDF")
    finally:
        db.close()
        Base.metadata.drop_all(engine)


def test_energy_excel_endpoints_are_registered_in_openapi() -> None:
    app = FastAPI()
    app.include_router(router)
    app.include_router(reports_router)
    schema = app.openapi()
    for path in [
        "/api/energy-schedule/consumption",
        "/api/energy-schedule/consumption/upload",
        "/api/energy-schedule/calculation-comparison",
        "/api/energy-schedule/calculation-trace",
        "/api/energy-schedule/calculation-results",
        "/api/calculate/energy-schedule",
        "/api/reports/energy-schedule/excel-calculation",
        "/api/reports/energy-schedule/pdf-calculation",
    ]:
        assert path in schema["paths"]


def pytest_approx(value: float):
    import pytest

    return pytest.approx(value, rel=0, abs=0.000001)


def sample_workbook_inputs() -> list[DailyCalculationInput]:
    return [
        DailyCalculationInput(
            portfolio_id=1,
            trading_date=date(2026, 8, 1),
            consumption=ConsumptionEntry(
                portfolio_id=1,
                consumption_date=date(2026, 8, 1),
                c1_kwh=1500,
                c2_kwh=9300,
                c4_kwh=18300,
                c5_kwh=5150,
                base_tariff_per_unit=7.5,
            ),
            markets={
                "gdam": MarketInput(
                    market="gdam",
                    purchased_mwh=2.55,
                    iex_payout=6850.77,
                    ctu_loss_pct=0.0379,
                    ctu_charge=122.94,
                    nldc_fee=5.64,
                    bucket_mwh={"c1": 2.55, "c2": 0.0, "c4": 0.0, "c5": 0.0},
                ),
                "dam": MarketInput(market="dam"),
                "rtm": MarketInput(market="rtm"),
            },
        ),
        DailyCalculationInput(
            portfolio_id=1,
            trading_date=date(2026, 8, 2),
            consumption=ConsumptionEntry(
                portfolio_id=1,
                consumption_date=date(2026, 8, 2),
                c1_kwh=2050,
                c2_kwh=10000,
                c4_kwh=19250,
                c5_kwh=3650,
                base_tariff_per_unit=8.5,
            ),
            markets={
                "gdam": MarketInput(
                    market="gdam",
                    purchased_mwh=2.75,
                    iex_payout=9989.92,
                    ctu_loss_pct=0.0379,
                    ctu_charge=122.94,
                    nldc_fee=6.75,
                    bucket_mwh={"c1": 2.75, "c2": 0.0, "c4": 0.0, "c5": 0.0},
                ),
                "dam": MarketInput(market="dam"),
                "rtm": MarketInput(market="rtm"),
            },
        ),
        DailyCalculationInput(
            portfolio_id=1,
            trading_date=date(2026, 8, 3),
            consumption=ConsumptionEntry(
                portfolio_id=1,
                consumption_date=date(2026, 8, 3),
                c1_kwh=1800,
                c2_kwh=10650,
                c4_kwh=20350,
                c5_kwh=3800,
                base_tariff_per_unit=8.5,
            ),
            markets={
                "gdam": MarketInput(
                    market="gdam",
                    purchased_mwh=2.60969,
                    iex_payout=10139.4156,
                    ctu_loss_pct=0.0379,
                    ctu_charge=122.94,
                    nldc_fee=7.07,
                    bucket_mwh={"c1": 2.60969, "c2": 0.0, "c4": 0.0, "c5": 0.0},
                ),
                "dam": MarketInput(market="dam"),
                "rtm": MarketInput(market="rtm"),
            },
        ),
    ]
