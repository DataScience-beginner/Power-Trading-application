"""Excel export builder for saved energy schedule calculation results."""

from __future__ import annotations

from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from database.models import MonthlyCalculation, Portfolio


TITLE_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FILL = PatternFill("solid", fgColor="D9EAF7")
GREEN_FILL = PatternFill("solid", fgColor="E2F0D9")
THIN_BORDER = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
    bottom=Side(style="thin", color="D9D9D9"),
)


def _clean_number(value: Any) -> float | None:
    return value if isinstance(value, (int, float)) else None


def _display_label(value: str) -> str:
    return value.replace("_", " ").title()


def _daily_output(row: dict[str, Any], mode: str) -> dict[str, Any]:
    if mode == "compare":
        return row.get("corrected") or {}
    return row.get(mode) or row


def _summary_parts(monthly_summary: MonthlyCalculation) -> dict[str, Any]:
    data = monthly_summary.calculation_data or {}
    savings = data.get("savings_sheet") or {}
    return {
        "data": data,
        "mode": data.get("calculation_mode") or "compare",
        "totals": savings.get("totals") or {},
        "completion": data.get("completion") or {},
        "slot_rows": (data.get("slot_wise_consolidate") or {}).get("rows") or [],
        "daily_results": data.get("daily_results") or [],
        "assumptions": data.get("assumptions") or {},
    }


def _style_title(ws, title: str, width: int) -> None:
    ws.cell(row=1, column=1, value=title)
    ws.cell(row=1, column=1).font = Font(size=16, bold=True, color="FFFFFF")
    ws.cell(row=1, column=1).fill = TITLE_FILL
    ws.cell(row=1, column=1).alignment = Alignment(horizontal="center")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=width)


def _style_table(ws, header_row: int, max_col: int) -> None:
    for cell in ws[header_row]:
        if cell.column <= max_col:
            cell.font = Font(bold=True)
            cell.fill = HEADER_FILL
            cell.alignment = Alignment(horizontal="center")
    for row in ws.iter_rows(min_row=header_row, max_row=ws.max_row, max_col=max_col):
        for cell in row:
            cell.border = THIN_BORDER
            if isinstance(cell.value, (int, float)):
                cell.number_format = '#,##0.00'


def _autofit(ws) -> None:
    for column_cells in ws.columns:
        letter = get_column_letter(column_cells[0].column)
        width = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column_cells)
        ws.column_dimensions[letter].width = min(max(width + 2, 12), 32)


def generate_energy_excel_calculation_report(
    monthly_summary: MonthlyCalculation,
    portfolio: Portfolio | None = None,
) -> BytesIO:
    """Generate a clean Excel report from a persisted monthly summary row."""
    parts = _summary_parts(monthly_summary)
    data = parts["data"]
    mode = parts["mode"]
    totals = parts["totals"]
    completion = parts["completion"]
    slot_rows = parts["slot_rows"]
    daily_results = parts["daily_results"]
    assumptions = parts["assumptions"]

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    _style_title(ws, "Energy Schedule Savings Report", 4)
    summary_rows = [
        ("Portfolio", portfolio.portfolio_code if portfolio else monthly_summary.portfolio_id),
        ("Portfolio Name", portfolio.portfolio_name if portfolio else ""),
        ("Period", f"{monthly_summary.month:02d}-{monthly_summary.year}"),
        ("Calculation Version", data.get("calculation_version")),
        ("Calculation Mode", mode),
        ("Completed Days", f"{completion.get('days_complete', 0)} / {completion.get('days_total', 0)}"),
        ("Missing Input Count", completion.get("missing_input_count", 0)),
        ("Total IEX Cost", totals.get("iex_price")),
        ("Equivalent EB Cost", totals.get("equivalent_eb_price")),
        ("Total Cost Saving", totals.get("total_cost_saving")),
        ("IEX Price Per Unit Avg", totals.get("iex_price_per_unit_average")),
        ("EB Price Per Unit Avg", totals.get("equivalent_eb_price_per_unit_average")),
        ("Saving Per Unit Avg", totals.get("cost_saving_per_unit_average")),
    ]
    for idx, (label, value) in enumerate(summary_rows, start=3):
        ws.cell(idx, 1, label)
        ws.cell(idx, 2, value)
    ws["A3"].fill = GREEN_FILL
    _style_table(ws, 3, 2)
    _autofit(ws)

    ws_daily = wb.create_sheet("Daily Results")
    _style_title(ws_daily, "Daily Results", 8)
    headers = ["Date", "Status", "IEX Cost", "IEX / Unit", "Equivalent EB Cost", "EB / Unit", "Cost Saving", "Saving / Unit"]
    ws_daily.append([])
    ws_daily.append(headers)
    for item in daily_results:
        output = _daily_output(item, mode)
        ws_daily.append([
            item.get("trading_date"),
            "Complete" if output.get("is_complete") else "Incomplete",
            _clean_number(output.get("b44_iex_price")),
            _clean_number(output.get("b45_iex_price_per_unit")),
            _clean_number(output.get("b46_eb_price")),
            _clean_number(output.get("b47_eb_price_per_unit")),
            _clean_number(output.get("cost_saving")),
            _clean_number(output.get("cost_saving_per_unit")),
        ])
    _style_table(ws_daily, 3, len(headers))
    _autofit(ws_daily)

    ws_slots = wb.create_sheet("Slot Wise Consolidate")
    _style_title(ws_slots, "Slot Wise Consolidate", 4)
    ws_slots.append([])
    ws_slots.append(["Slot", "Consumption kWh", "IEX Delivered kWh", "Balance kWh"])
    for row in slot_rows:
        ws_slots.append([
            row.get("slot"),
            _clean_number(row.get("consumption_kwh")),
            _clean_number(row.get("iex_delivered_kwh")),
            _clean_number(row.get("balance_kwh")),
        ])
    _style_table(ws_slots, 3, 4)
    _autofit(ws_slots)

    ws_assumptions = wb.create_sheet("Assumptions")
    _style_title(ws_assumptions, "Controlled Assumptions", 2)
    ws_assumptions.append([])
    ws_assumptions.append(["Assumption", "Value"])
    for key, value in assumptions.items():
        ws_assumptions.append([_display_label(key), value])
    _style_table(ws_assumptions, 3, 2)
    _autofit(ws_assumptions)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def generate_energy_pdf_calculation_report(
    monthly_summary: MonthlyCalculation,
    portfolio: Portfolio | None = None,
) -> BytesIO:
    """Generate a clean PDF summary from a persisted monthly summary row."""
    parts = _summary_parts(monthly_summary)
    totals = parts["totals"]
    completion = parts["completion"]
    mode = parts["mode"]
    data = parts["data"]
    slot_rows = parts["slot_rows"]

    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36)
    styles = getSampleStyleSheet()
    elements = [
        Paragraph("Energy Schedule Savings Report", styles["Title"]),
        Paragraph(f"Portfolio: {portfolio.portfolio_code if portfolio else monthly_summary.portfolio_id}", styles["Normal"]),
        Paragraph(f"Period: {monthly_summary.month:02d}-{monthly_summary.year}", styles["Normal"]),
        Spacer(1, 12),
    ]

    summary_data = [
        ["Metric", "Value"],
        ["Calculation Version", data.get("calculation_version") or "-"],
        ["Calculation Mode", mode],
        ["Completed Days", f"{completion.get('days_complete', 0)} / {completion.get('days_total', 0)}"],
        ["Missing Input Count", completion.get("missing_input_count", 0)],
        ["Total IEX Cost", f"{_clean_number(totals.get('iex_price')) or 0:,.2f}"],
        ["Equivalent EB Cost", f"{_clean_number(totals.get('equivalent_eb_price')) or 0:,.2f}"],
        ["Total Cost Saving", f"{_clean_number(totals.get('total_cost_saving')) or 0:,.2f}"],
        ["Saving Per Unit Avg", f"{_clean_number(totals.get('cost_saving_per_unit_average')) or 0:,.4f}"],
    ]
    summary_table = Table(summary_data, colWidths=[210, 260])
    summary_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E78")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("BACKGROUND", (0, 1), (-1, -1), colors.HexColor("#F7FBFF")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]))
    elements.extend([summary_table, Spacer(1, 18), Paragraph("Slot Wise Consolidate", styles["Heading2"])])

    slot_data = [["Slot", "Consumption kWh", "IEX Delivered kWh", "Balance kWh"]]
    for row in slot_rows:
        slot_data.append([
            row.get("slot"),
            f"{_clean_number(row.get('consumption_kwh')) or 0:,.2f}",
            f"{_clean_number(row.get('iex_delivered_kwh')) or 0:,.2f}",
            f"{_clean_number(row.get('balance_kwh')) or 0:,.2f}",
        ])
    slot_table = Table(slot_data, colWidths=[80, 130, 130, 130])
    slot_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#D9EAF7")),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("ALIGN", (1, 1), (-1, -1), "RIGHT"),
    ]))
    elements.append(slot_table)

    doc.build(elements)
    buffer.seek(0)
    return buffer
