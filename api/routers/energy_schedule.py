"""Core energy schedule routes."""

from datetime import date as dt_date
from datetime import datetime
from typing import Any, Optional

from fastapi import APIRouter, Body, Depends, File, HTTPException, UploadFile
from sqlalchemy.orm import Session

from api.schemas.energy_excel_calculation import (
    CalculationTraceResponse,
    CalculationRunRequest,
    CalculationRunResponse,
    ComparisonResponse,
    ConsumptionEntry,
    ConsumptionResponse,
    ConsumptionSaveRequest,
    SavedCalculationResultsResponse,
    SavedCalculationRow,
)
from api.services.energy_excel_calculation_service import (
    CalculationAssumptions,
    CALCULATION_VERSION,
    build_calculation_inputs,
    build_calculation_trace,
    calculate_savings_sheet,
    calculate_slot_wise_consolidate,
    compare_daily,
    list_consumption,
    run_calculation,
    upsert_consumption,
)
from api.services.energy_schedule_workflow_service import (
    build_workflow_status,
    rebuild_workflow_energy_schedule,
    seed_workflow_mock_data,
)
from api.security.chat_auth import require_admin
from database.config import get_db
from database.energy_schedule_crud import get_all_daily_entries, get_all_month_sheets
from database.energy_schedule_service import calculator
from database.models import EnergyScheduleDay, EnergyScheduleMonth, MonthlyCalculation


router = APIRouter(tags=["energy-schedule"])


def _saved_calculation_row(row: MonthlyCalculation) -> SavedCalculationRow:
    """Map a persisted calculation row to its API response shape."""
    return SavedCalculationRow(
        id=row.id,
        calculation_date=row.calculation_date,
        day=row.day,
        calculation_type=row.calculation_type,
        calculated_at=row.calculated_at,
        updated_at=row.updated_at,
        total_scheduled_mwh=float(row.total_scheduled_mwh or 0.0),
        total_cost=float(row.total_cost or 0.0),
        net_profit_loss=float(row.net_profit_loss or 0.0),
        calculation_data=row.calculation_data,
    )


def _parse_date(value: str) -> dt_date:
    """Parse API date strings consistently."""
    return datetime.fromisoformat(value.split("T")[0]).date()


@router.get(
    "/api/energy-schedule/days",
    response_model=dict[str, Any],
    summary="List energy schedule days",
    description="Returns daily energy schedule records with completeness flags, CTU losses, savings, costs, and consumption.",
)
async def get_energy_schedule_days(
    portfolio_id: Optional[int] = None,
    start_date: Optional[str] = None,
    end_date: Optional[str] = None,
    complete_only: bool = False,
    db: Session = Depends(get_db),
) -> dict[str, Any]:
    """Return energy schedule day records with calculations."""
    try:
        query = db.query(EnergyScheduleDay).join(EnergyScheduleMonth)

        if portfolio_id:
            query = query.filter(EnergyScheduleMonth.portfolio_id == portfolio_id)
        if complete_only:
            query = query.filter(EnergyScheduleDay.is_complete == 1)
        if start_date:
            start = datetime.strptime(start_date, "%Y-%m-%d").date()
            query = query.filter(EnergyScheduleDay.trading_date >= start)
        if end_date:
            end = datetime.strptime(end_date, "%Y-%m-%d").date()
            query = query.filter(EnergyScheduleDay.trading_date <= end)

        days = query.order_by(EnergyScheduleDay.trading_date).all()

        result = [
            {
                "id": day.id,
                "trading_date": str(day.trading_date),
                "is_complete": bool(day.is_complete),
                "has_gdam": bool(day.has_gdam_data),
                "has_dam": bool(day.has_dam_data),
                "has_rtm": bool(day.has_rtm_data),
                "has_sch": bool(day.has_sch_data),
                "total_scheduled_mwh": round(day.total_scheduled_mwh, 2) if day.total_scheduled_mwh else 0,
                "ctu_losses_mwh": round(day.ctu_losses_mwh, 2) if day.ctu_losses_mwh else 0,
                "ctu_losses_percent": round(day.ctu_losses_percent, 2) if day.ctu_losses_percent else 0,
                "energy_savings_mwh": round(day.energy_savings_mwh, 2) if day.energy_savings_mwh else 0,
                "total_cost": round(day.total_cost, 2) if day.total_cost else 0,
                "total_consumption_mwh": (
                    round(day.total_consumption_after_losses_mwh, 2)
                    if day.total_consumption_after_losses_mwh
                    else 0
                ),
            }
            for day in days
        ]

        return {"success": True, "count": len(result), "days": result}

    except HTTPException:
        raise
    except Exception as e:
        import traceback

        print(traceback.format_exc())
        return {"success": False, "error": str(e)}


@router.get(
    "/api/energy-schedule/consumption",
    response_model=ConsumptionResponse,
    summary="List energy schedule consumption inputs",
    description="Returns client-provided C1, C2, C4, and C5 consumption inputs used by the Excel-conversion calculator.",
)
async def get_energy_schedule_consumption(
    portfolio_id: int,
    start_date: Optional[str] = None,
    end_date: Optional[str] = None,
    db: Session = Depends(get_db),
) -> ConsumptionResponse:
    """Read manual or uploaded consumption rows for a portfolio and date range."""
    records = list_consumption(
        db,
        portfolio_id=portfolio_id,
        start_date=_parse_date(start_date) if start_date else None,
        end_date=_parse_date(end_date) if end_date else None,
    )
    return ConsumptionResponse(success=True, count=len(records), records=records)


@router.post(
    "/api/energy-schedule/consumption",
    response_model=ConsumptionResponse,
    summary="Save energy schedule consumption inputs",
    description="Creates or updates client-provided C1, C2, C4, and C5 consumption rows for Excel-conversion calculations.",
)
async def save_energy_schedule_consumption(
    request: ConsumptionSaveRequest,
    db: Session = Depends(get_db),
) -> ConsumptionResponse:
    """Create or update manual consumption rows."""
    records = upsert_consumption(db, request.records)
    return ConsumptionResponse(success=True, count=len(records), records=records)


@router.post(
    "/api/energy-schedule/consumption/upload",
    response_model=ConsumptionResponse,
    summary="Upload energy schedule consumption inputs",
    description="Uploads a monthly consumption workbook or CSV-style file with Date, C1, C2, C4, C5, and optional tariff columns.",
)
async def upload_energy_schedule_consumption(
    portfolio_id: int,
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
) -> ConsumptionResponse:
    """Upload client consumption details from a workbook-like file."""
    try:
        import csv
        from io import BytesIO, StringIO

        import openpyxl

        content = await file.read()
        suffix = (file.filename or "").lower()
        rows: list[list[Any]] = []

        if suffix.endswith(".csv"):
            rows = list(csv.reader(StringIO(content.decode("utf-8-sig"))))
        elif suffix.endswith((".xlsx", ".xlsm")):
            workbook = openpyxl.load_workbook(BytesIO(content), data_only=True)
            sheet = workbook["Consumption Details"] if "Consumption Details" in workbook.sheetnames else workbook.active
            rows = [[cell for cell in row] for row in sheet.iter_rows(values_only=True)]
        else:
            raise HTTPException(status_code=400, detail="Upload .xlsx, .xlsm, or .csv consumption files only.")

        records: list[ConsumptionEntry] = []
        for raw in rows:
            if not raw or raw[0] in (None, "", "Date"):
                continue
            try:
                item_date = raw[0].date() if hasattr(raw[0], "date") else _parse_date(str(raw[0]))
            except Exception:
                continue
            records.append(
                ConsumptionEntry(
                    portfolio_id=portfolio_id,
                    consumption_date=item_date,
                    c1_kwh=float(raw[1] or 0) if len(raw) > 1 else 0.0,
                    c2_kwh=float(raw[2] or 0) if len(raw) > 2 else 0.0,
                    c4_kwh=float(raw[3] or 0) if len(raw) > 3 else 0.0,
                    c5_kwh=float(raw[4] or 0) if len(raw) > 4 else 0.0,
                    base_tariff_per_unit=float(raw[5]) if len(raw) > 5 and raw[5] not in (None, "") else None,
                    source="upload",
                    notes=file.filename,
                )
            )

        if not records:
            raise HTTPException(status_code=400, detail="No valid consumption rows were found.")

        saved = upsert_consumption(db, records)
        return ConsumptionResponse(success=True, count=len(saved), records=saved)
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Could not parse consumption upload: {str(e)}") from e


@router.get(
    "/api/energy-schedule/workflow-status",
    response_model=dict[str, Any],
    summary="Get Energy Schedule workflow status",
    description="Returns admin workflow readiness across upload, Energy Schedule validation, consumption, calculation, and reports.",
)
async def get_energy_schedule_workflow_status(
    portfolio_id: int,
    year: int,
    month: int,
    mode: str = "compare",
    db: Session = Depends(get_db),
) -> dict[str, Any]:
    """Return a guided admin workflow status for one portfolio month."""
    if mode not in {"parity", "corrected", "compare"}:
        raise HTTPException(status_code=400, detail="mode must be parity, corrected, or compare")
    return build_workflow_status(db, portfolio_id=portfolio_id, year=year, month=month, mode=mode)


@router.post(
    "/api/energy-schedule/rebuild",
    response_model=dict[str, Any],
    summary="Rebuild Energy Schedule from uploads",
    description="Builds or refreshes Energy Schedule rows from uploaded DOR and SCH data for the selected portfolio month.",
)
async def rebuild_energy_schedule_workflow(
    request_data: dict[str, Any] = Body(...),
    db: Session = Depends(get_db),
) -> dict[str, Any]:
    """Refresh Energy Schedule rows for every uploaded date in a month."""
    try:
        portfolio_id = int(request_data["portfolio_id"])
        year = int(request_data["year"])
        month = int(request_data["month"])
    except (KeyError, TypeError, ValueError) as exc:
        raise HTTPException(status_code=400, detail="portfolio_id, year, and month are required") from exc
    return rebuild_workflow_energy_schedule(db, portfolio_id=portfolio_id, year=year, month=month)


@router.post(
    "/api/energy-schedule/workflow-demo-seed",
    response_model=dict[str, Any],
    summary="Seed Energy Schedule workflow demo data",
    description="Creates controlled synthetic upload, Energy Schedule, and consumption rows for admin workflow testing.",
)
async def seed_energy_schedule_workflow_demo(
    request_data: dict[str, Any] = Body(...),
    db: Session = Depends(get_db),
    _admin: Any = Depends(require_admin),
) -> dict[str, Any]:
    """Seed a small synthetic dataset for the guided workflow page."""
    try:
        portfolio_id = int(request_data["portfolio_id"])
        year = int(request_data["year"])
        month = int(request_data["month"])
        days = int(request_data.get("days", 3))
    except (KeyError, TypeError, ValueError) as exc:
        raise HTTPException(status_code=400, detail="portfolio_id, year, month, and optional numeric days are required") from exc
    if days < 1 or days > 31:
        raise HTTPException(status_code=400, detail="days must be between 1 and 31")
    return seed_workflow_mock_data(db, portfolio_id=portfolio_id, year=year, month=month, days=days)


@router.post(
    "/api/calculate/energy-schedule",
    response_model=dict[str, Any],
    summary="Calculate energy schedule",
    description="Runs the energy schedule calculator for a requested calculation date.",
)
async def calculate_energy_schedule(
    request_data: dict[str, Any] = Body(...),
    db: Session = Depends(get_db),
) -> dict[str, Any]:
    """Calculate energy schedule using legacy or Excel-conversion logic."""
    try:
        if request_data.get("portfolio_id"):
            request = CalculationRunRequest(**request_data)
            if request.calculation_date:
                year = request.calculation_date.year
                month = request.calculation_date.month
                day = request.calculation_date.day
            else:
                if request.year is None or request.month is None:
                    raise HTTPException(
                        status_code=400,
                        detail="year and month are required unless calculation_date is provided.",
                    )
                year = request.year
                month = request.month
                day = request.day

            results = run_calculation(
                db,
                portfolio_id=request.portfolio_id,
                year=year,
                month=month,
                day=day,
                mode=request.mode,
            )
            response = CalculationRunResponse(
                success=True,
                mode=request.mode,
                portfolio_id=request.portfolio_id,
                year=year,
                month=month,
                day=day,
                days_processed=len(results),
                results=results,
            )
            return response.model_dump(mode="json")

        calculation_date_str = request_data.get("calculation_date")
        print(f"🔍 Calculate request - calculation_date: {calculation_date_str}")

        if calculation_date_str:
            try:
                calc_date = datetime.fromisoformat(calculation_date_str).date()
            except Exception:
                calc_date = datetime.strptime(calculation_date_str.split("T")[0], "%Y-%m-%d").date()
        else:
            calc_date = dt_date.today()

        print(f"📅 Using calculation date: {calc_date}")
        return calculator.calculate_energy_schedule(calc_date, db)

    except HTTPException:
        raise
    except Exception as e:
        import traceback

        error_details = traceback.format_exc()
        print(f"❌ Error in calculate_energy_schedule: {error_details}")

        return {
            "success": False,
            "message": f"Calculation failed: {str(e)}",
            "error": str(e),
        }


@router.get(
    "/api/energy-schedule/calculation-comparison",
    response_model=ComparisonResponse,
    summary="Compare Excel-parity and corrected energy calculations",
    description="Returns day-wise Excel-parity and corrected backend outputs for manual QA of the new calculation engine.",
)
async def get_energy_schedule_calculation_comparison(
    portfolio_id: int,
    year: int,
    month: int,
    day: Optional[int] = None,
    db: Session = Depends(get_db),
) -> ComparisonResponse:
    """Compare parity and corrected calculations for a day or month."""
    import calendar

    assumptions = CalculationAssumptions()
    dates = (
        [dt_date(year, month, day)]
        if day is not None
        else [dt_date(year, month, item) for item in range(1, calendar.monthrange(year, month)[1] + 1)]
    )
    daily_inputs = [build_calculation_inputs(db, portfolio_id, trading_date) for trading_date in dates]
    comparisons = [compare_daily(inputs, assumptions) for inputs in daily_inputs]
    return ComparisonResponse(
        success=True,
        portfolio_id=portfolio_id,
        year=year,
        month=month,
        count=len(comparisons),
        tolerance=assumptions.tolerance,
        comparisons=comparisons,
        savings_sheet=calculate_savings_sheet(comparisons),
        slot_wise_consolidate=calculate_slot_wise_consolidate(daily_inputs),
    )


@router.get(
    "/api/energy-schedule/calculation-trace",
    response_model=CalculationTraceResponse,
    summary="Trace energy schedule calculation inputs",
    description="Returns an internal QA trace from uploaded files and DB energy schedule fields to daily and monthly Excel-conversion outputs.",
)
async def get_energy_schedule_calculation_trace(
    portfolio_id: int,
    year: int,
    month: int,
    day: Optional[int] = None,
    db: Session = Depends(get_db),
) -> CalculationTraceResponse:
    """Return a compact internal QA trace for Excel-conversion calculations."""
    trace = build_calculation_trace(db, portfolio_id=portfolio_id, year=year, month=month, day=day)
    return CalculationTraceResponse(
        success=True,
        portfolio_id=portfolio_id,
        year=year,
        month=month,
        day=day,
        count=len(trace["daily_outputs"]),
        trace=trace,
    )


@router.get(
    "/api/energy-schedule/calculation-results",
    response_model=SavedCalculationResultsResponse,
    summary="Get saved Excel-conversion calculation results",
    description="Returns persisted Excel-conversion v1 calculation rows without recalculating.",
)
async def get_energy_schedule_calculation_results(
    portfolio_id: int,
    year: int,
    month: int,
    mode: str = "compare",
    day: Optional[int] = None,
    db: Session = Depends(get_db),
) -> SavedCalculationResultsResponse:
    """Return saved calculation rows for the admin review/export workflow."""
    if mode not in {"parity", "corrected", "compare"}:
        raise HTTPException(status_code=400, detail="mode must be parity, corrected, or compare")

    query = db.query(MonthlyCalculation).filter(
        MonthlyCalculation.portfolio_id == portfolio_id,
        MonthlyCalculation.year == year,
        MonthlyCalculation.month == month,
        MonthlyCalculation.calculation_type == f"{CALCULATION_VERSION}:{mode}",
    )
    if day is not None:
        query = query.filter(MonthlyCalculation.day == day)

    rows = query.order_by(MonthlyCalculation.calculation_date.asc()).all()
    monthly_summary = db.query(MonthlyCalculation).filter(
        MonthlyCalculation.portfolio_id == portfolio_id,
        MonthlyCalculation.year == year,
        MonthlyCalculation.month == month,
        MonthlyCalculation.calculation_type == f"{CALCULATION_VERSION}:monthly:{mode}",
    ).first()
    all_calculated_at = [row.calculated_at for row in rows if row.calculated_at]
    if monthly_summary and monthly_summary.calculated_at:
        all_calculated_at.append(monthly_summary.calculated_at)
    latest_calculated_at = max(all_calculated_at, default=None)

    return SavedCalculationResultsResponse(
        success=True,
        portfolio_id=portfolio_id,
        year=year,
        month=month,
        mode=mode,
        day=day,
        count=len(rows),
        latest_calculated_at=latest_calculated_at,
        monthly_summary=_saved_calculation_row(monthly_summary) if monthly_summary else None,
        results=[_saved_calculation_row(row) for row in rows],
    )


@router.get(
    "/api/energy-schedule/status",
    response_model=dict[str, Any],
    summary="Get energy schedule status",
    description="Returns recent energy schedule calculation status records.",
)
async def get_energy_schedule_status(
    start_date: Optional[str] = None,
    end_date: Optional[str] = None,
    db: Session = Depends(get_db),
) -> dict[str, Any]:
    """Return recent energy schedule calculation status."""
    try:
        query = db.query(EnergyScheduleDay).order_by(EnergyScheduleDay.trading_date.desc()).limit(30)
        schedules_objs = query.all()

        schedules = [
            {
                "id": s.id,
                "month_sheet_id": s.month_sheet_id,
                "trading_date": str(s.trading_date),
                "scheduled_mwh": float(s.gdam_scheduled_quantity_mwh or 0)
                + float(s.dam_scheduled_quantity_mwh or 0)
                + float(s.rtm_scheduled_quantity_mwh or 0),
                "consumption_after_losses_mwh": float(s.sch_consumption_after_losses_mwh or 0),
                "gdam_cost": float(s.gdam_cost) if s.gdam_cost else 0.0,
                "dam_cost": float(s.dam_cost) if s.dam_cost else 0.0,
                "rtm_cost": float(s.rtm_cost) if s.rtm_cost else 0.0,
                "ctu_loss_percentage": float(s.ctu_loss_percentage) if s.ctu_loss_percentage else 0.0,
            }
            for s in schedules_objs
        ]

        return {"success": True, "count": len(schedules), "schedules": schedules}

    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Error fetching energy schedule status: {str(e)}") from e


@router.get(
    "/api/energy-schedule/months",
    response_model=list[dict[str, Any]],
    summary="List energy schedule months",
    description="Returns month sheets with summary metrics for the Energy Schedule frontend.",
)
async def get_energy_schedule_months(
    portfolio_id: Optional[int] = None,
    db: Session = Depends(get_db),
) -> list[dict[str, Any]]:
    """Return all energy schedule month sheets, optionally filtered by portfolio."""
    try:
        month_sheets = get_all_month_sheets(db, portfolio_id)

        result = []
        for sheet in month_sheets:
            total_nldc = 0.0
            total_ctu = 0.0
            for day in sheet.daily_entries:
                total_nldc += (day.gdam_nldc_fee or 0) + (day.dam_nldc_fee or 0) + (day.rtm_nldc_fee or 0)
                total_ctu += (day.gdam_ctu_charges or 0) + (day.dam_ctu_charges or 0) + (day.rtm_ctu_charges or 0)

            result.append(
                {
                    "id": sheet.id,
                    "portfolio_id": sheet.portfolio_id,
                    "year": sheet.year,
                    "month": sheet.month,
                    "month_name": sheet.month_name,
                    "total_scheduled_mwh": sheet.total_scheduled_mwh,
                    "total_consumption_after_losses_mwh": sheet.total_consumption_after_losses_mwh,
                    "total_energy_savings": sheet.total_energy_savings_mwh,
                    "total_gdam_cost": sheet.total_gdam_cost,
                    "total_dam_cost": sheet.total_dam_cost,
                    "total_rtm_cost": sheet.total_rtm_cost,
                    "total_nldc_fees": total_nldc,
                    "total_ctu_charges": total_ctu,
                    "total_cost": sheet.total_gdam_cost + sheet.total_dam_cost + sheet.total_rtm_cost,
                    "average_ctu_losses": sheet.average_ctu_losses_percent,
                    "total_days_completed": sheet.days_completed,
                    "is_complete": sheet.is_complete,
                    "created_at": sheet.created_at.isoformat() if sheet.created_at else None,
                    "updated_at": sheet.updated_at.isoformat() if sheet.updated_at else None,
                }
            )

        return result

    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Error fetching month sheets: {str(e)}") from e


@router.get(
    "/api/energy-schedule/days",
    response_model=dict[str, Any],
    summary="List energy schedule days by date",
    description="Legacy duplicate route retained in the same order for compatibility; current FastAPI behavior uses the first matching route.",
)
async def get_energy_schedule_days_by_date(
    portfolio_id: Optional[int] = None,
    start_date: Optional[str] = None,
    end_date: Optional[str] = None,
    db: Session = Depends(get_db),
) -> dict[str, Any]:
    """Return daily entries by date range. Kept for compatibility with current route registration order."""
    try:
        query = db.query(EnergyScheduleDay).join(EnergyScheduleMonth)

        if portfolio_id:
            query = query.filter(EnergyScheduleMonth.portfolio_id == portfolio_id)
        if start_date:
            start = datetime.fromisoformat(start_date).date()
            query = query.filter(EnergyScheduleDay.trading_date >= start)
        if end_date:
            end = datetime.fromisoformat(end_date).date()
            query = query.filter(EnergyScheduleDay.trading_date <= end)

        daily_entries = query.order_by(EnergyScheduleDay.trading_date).all()

        result = [
            {
                "id": entry.id,
                "trading_date": entry.trading_date.isoformat(),
                "portfolio_id": entry.month_sheet.portfolio_id,
                "total_scheduled_mwh": float(entry.total_scheduled_mwh or 0),
                "total_consumption_after_losses_mwh": float(entry.total_consumption_after_losses_mwh or 0),
                "ctu_losses_percent": float(entry.ctu_losses_percent or 0),
                "ctu_losses_mwh": float(entry.ctu_losses_mwh or 0),
                "gdam_cost": float(entry.gdam_cost or 0),
                "dam_cost": float(entry.dam_cost or 0),
                "rtm_cost": float(entry.rtm_cost or 0),
                "total_cost": float(entry.total_cost or 0),
                "energy_savings_mwh": float(entry.energy_savings_mwh or 0),
                "total_nldc_fee": float(entry.total_nldc_fee or 0),
                "total_ctu_charges": float(entry.total_ctu_charges or 0),
                "has_gdam_data": bool(entry.has_gdam_data),
                "has_dam_data": bool(entry.has_dam_data),
                "has_rtm_data": bool(entry.has_rtm_data),
                "has_sch_data": bool(entry.has_sch_data),
                "is_complete": bool(entry.is_complete),
                "calculated_at": entry.calculated_at.isoformat() if entry.calculated_at else None,
            }
            for entry in daily_entries
        ]

        return {"success": True, "count": len(result), "days": result}

    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Error fetching daily entries: {str(e)}") from e


@router.get(
    "/api/energy-schedule/months/{month_sheet_id}/days",
    response_model=dict[str, Any],
    summary="List days for energy schedule month",
    description="Returns all daily entries and DOR/SCH calculation fields for a month sheet.",
)
async def get_energy_schedule_days_by_month(
    month_sheet_id: int,
    db: Session = Depends(get_db),
) -> dict[str, Any]:
    """Return all daily entries for a month sheet."""
    try:
        daily_entries = get_all_daily_entries(db, month_sheet_id)

        result = []
        for entry in daily_entries:
            result.append(
                {
                    "id": entry.id,
                    "trading_date": entry.trading_date.isoformat(),
                    "day_of_month": entry.day_of_month,
                    "gdam": {
                        "nldc_fee": entry.gdam_nldc_fee,
                        "ctu_charges": entry.gdam_ctu_charges,
                        "cost": entry.gdam_cost,
                        "scheduled_mwh": entry.gdam_scheduled_quantity_mwh,
                    },
                    "dam": {
                        "nldc_fee": entry.dam_nldc_fee,
                        "ctu_charges": entry.dam_ctu_charges,
                        "cost": entry.dam_cost,
                        "scheduled_mwh": entry.dam_scheduled_quantity_mwh,
                    },
                    "rtm": {
                        "nldc_fee": entry.rtm_nldc_fee,
                        "ctu_charges": entry.rtm_ctu_charges,
                        "cost": entry.rtm_cost,
                        "scheduled_mwh": entry.rtm_scheduled_quantity_mwh,
                    },
                    "consumption_after_losses_mwh": entry.total_consumption_after_losses_mwh,
                    "regional_loss_percent": entry.regional_loss_percent,
                    "state_loss_percent": entry.state_loss_percent,
                    "combined_loss_percent": entry.combined_loss_percent,
                    "total_scheduled_mwh": entry.total_scheduled_mwh,
                    "ctu_losses_percent": entry.ctu_losses_percent,
                    "ctu_losses_mwh": entry.ctu_losses_mwh,
                    "energy_savings_mwh": entry.energy_savings_mwh,
                    "total_nldc_fee": entry.total_nldc_fee,
                    "total_ctu_charges": entry.total_ctu_charges,
                    "total_cost": entry.total_cost,
                    "is_complete": entry.is_complete,
                    "has_gdam_data": entry.has_gdam_data,
                    "has_dam_data": entry.has_dam_data,
                    "has_rtm_data": entry.has_rtm_data,
                    "has_sch_data": entry.has_sch_data,
                    "calculated_at": entry.calculated_at.isoformat() if entry.calculated_at else None,
                }
            )

        return {"success": True, "count": len(result), "daily_entries": result}

    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Error fetching daily entries: {str(e)}") from e
